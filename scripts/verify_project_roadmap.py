"""检查项目过程Word与渲染产物；依赖requirements-documents.txt。"""

from datetime import datetime, timezone
import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "output/deliverables/MathorCup项目全过程与当前进度.docx"
QA = ROOT / "output/document_qa/project_roadmap"


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--render-backend", default="not specified; consult renderer log")
    parser.add_argument(
        "--visual-qa-confirmed",
        action="store_true",
        help="Record manual page inspection only after it has actually been performed",
    )
    args = parser.parse_args()
    doc = Document(DOCX)
    text = "\n".join(p.text for p in doc.paragraphs)
    pdf = QA / "project-roadmap.pdf"
    pages = PdfReader(pdf).pages
    evidence = json.loads(
        (ROOT / "output/variable_relations/verification.json").read_text(encoding="utf-8")
    )
    refs = [
        "docs/environment-setup.md",
        "docs/git-usage.md",
        "docs/data-audit.md",
        "docs/spatial-audit.md",
        "docs/metadata-review.md",
        "docs/observation-exploration.md",
        "docs/variable-relations.md",
        "docs/pending-questions.md",
        "config/analysis-constraints.yml",
    ]
    with ZipFile(DOCX) as archive:
        package_valid = archive.testzip() is None
    checks = {
        "valid_docx_package": package_valid,
        "six_rendered_pages": len(pages) == 6
        and all((QA / f"page-{i + 1}.png").is_file() for i in range(6)),
        "ten_project_stages": len(doc.tables[0].rows) == 11,
        "current_stage_explicitly_first_question_preliminary": "当前位于第一问的初步特征解释与变量关系探索阶段"
        in text
        and "不能把第一问整体标记为已完成" in text,
        "event_stage_and_rain_definitions_still_pending": "不能自动视作同一时刻" in text
        and "降水累计、前期湿润指标和最终灾害时间标签继续暂缓" in text,
        "title_style_without_paragraph_border": not list(
            doc.styles["Title"]._element.iter(qn("w:pBdr"))
        ),
        "title_and_heading_styles_black": all(
            str(doc.styles[name].font.color.rgb) == "000000"
            for name in ["Title", "Subtitle", "Heading 1", "Heading 2"]
        ),
        "table_headers_repeat_and_rows_not_split": all(
            t.rows[0]._tr.find(".//" + qn("w:tblHeader")) is not None
            and all(r._tr.find(".//" + qn("w:cantSplit")) is not None for r in t.rows)
            for t in doc.tables
        ),
        "all_local_source_references_exist": all((ROOT / p).is_file() for p in refs),
        "analysis_verified_and_constraints_unchanged": evidence["passed"]
        and sha256(ROOT / "config/analysis-constraints.yml") == evidence["constraints_sha256"],
        "all_eight_raw_archives_unchanged": all(
            sha256(ROOT / "data/附件" / name) == value
            for name, value in evidence["source_hashes"].items()
        ),
        "no_unresolved_template_placeholders": all(
            token not in text for token in ["TODO", "TBD", "[[", "{{"]
        ),
    }
    result = {
        "passed": all(checks.values()),
        "checked_at_utc": datetime.now(timezone.utc).isoformat(),
        "checks": checks,
        "docx_sha256": sha256(DOCX),
        "pdf_sha256": sha256(pdf),
        "render_backend": args.render_backend,
        "raster_backend": "pypdfium2",
        "pages": len(pages),
        "visual_qa": "Manual page inspection confirmed by caller"
        if args.visual_qa_confirmed
        else "Not confirmed; automated checks cover structure only",
    }
    (QA / "verification.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    for key, value in checks.items():
        print(f"{'PASS' if value else 'FAIL'} {key}")
    if not result["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
