"""安装requirements-documents.txt中的可选依赖后生成项目过程Word。

这份过程说明不是正式参赛论文模板；原始数据和分析约束保持只读。
"""
from pathlib import Path
import json

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/deliverables/MathorCup项目全过程与当前进度.docx'
BLACK = RGBColor(0,0,0)


def font(run, size=11, bold=False):
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')


def paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        font(p.add_run(bold_lead),bold=True)
        font(p.add_run(text[len(bold_lead):]))
    else:
        font(p.add_run(text))
    return p


def heading(doc, text, level=1):
    return doc.add_heading(text,level=level)


def grid(doc, headers, rows, widths, highlight=None):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge=OxmlElement('w:'+side)
        for key,value in [('val','single'),('sz','5'),('color','D9D9D9')]:
            edge.set(qn('w:'+key),value)
        borders.append(edge)
    props.append(borders)
    for i,label in enumerate(headers):
        t.rows[0].cells[i].text=label
    repeated=OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeated)
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):
            cells[i].text=str(value)
    for r,row in enumerate(t.rows):
        tr_pr=row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement('w:cantSplit'))
        for c,cell in enumerate(row.cells):
            cell.width=Inches(widths[c])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd')
            shade.set(qn('w:fill'),'DBE8EF' if r==0 else 'EDF2F5' if highlight is not None and r==highlight+1 else 'FFFFFF')
            tcpr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                child=OxmlElement('w:'+edge)
                child.set(qn('w:w'),'85' if edge in ['top','bottom'] else '105')
                child.set(qn('w:type'),'dxa')
                margins.append(child)
            tcpr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.12
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if c==0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    font(run,size=10.5,bold=r==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t


def next_page(doc, title):
    doc.add_page_break()
    heading(doc,title)


def main():
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Inches(8.5)
    sec.page_height=Inches(11)
    sec.top_margin=Inches(.72)
    sec.bottom_margin=Inches(.68)
    sec.left_margin=Inches(.82)
    sec.right_margin=Inches(.82)
    sec.footer_distance=Inches(.3)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        style=doc.styles[name]
        style.font.name='Calibri'
        style.font.color.rgb=BLACK
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑' if name!='Normal' else '宋体')
        style.paragraph_format.space_after=Pt(7)
        style.paragraph_format.line_spacing=1.16
    doc.styles['Normal'].font.size=Pt(11)
    for name,size in [('Title',22),('Subtitle',11),('Heading 1',16),('Heading 2',12)]:
        doc.styles[name].font.size=Pt(size)
    doc.styles['Subtitle'].font.italic=False
    # The bundled default Title style may carry a blue bottom border.
    # Remove inherited title rules explicitly before visual rendering.
    for border in list(doc.styles['Title']._element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(9)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(7)
    for name in ['Heading 1','Heading 2']:
        doc.styles[name].paragraph_format.keep_with_next=True
    doc.core_properties.title='MathorCup项目全过程与当前进度'
    doc.core_properties.subject='赛道B初赛项目路线 证据边界与阶段产出'
    doc.core_properties.author='HCX'
    doc.core_properties.keywords='MathorCup 项目过程 当前进度 可复现'
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run('第 '),9)
    for text in ['PAGE','NUMPAGES']:
        f=OxmlElement('w:fldSimple')
        f.set(qn('w:instr'),text)
        footer._p.append(f)
        font(footer.add_run(' 页 / 共 ' if text=='PAGE' else ' 页'),9)

    # 第1页：当前位置及全过程索引。
    doc.add_paragraph('MathorCup项目全过程与当前进度','Title')
    doc.add_paragraph('赛道B初赛  吉隆泥石流项目  2026年10月27日','Subtitle')
    paragraph(doc,'当前位于第一问的初步特征解释与变量关系探索阶段。工具配置、文件审计、公开元数据复核和无标签质量探索已完成本轮工作；尚未进入最终触发评估、预报区间验证和灾害过程仿真。',bold_lead='当前位于第一问的初步特征解释与变量关系探索阶段。')
    paragraph(doc,'本说明用于安排工作和定位进度。各阶段按依赖关系推进，论文素材与复现记录同步积累；阶段序号不等于完成比例，不能据此认定项目已完成某个百分比。')
    heading(doc,'全过程与当前状态')
    grid(doc,['阶段','状态','主要产出'],[
        ('1 工具与工程','已完成本轮','Python建模环境 Git版本管理 TeX Live及QGIS'),
        ('2 数据与定义','核查已完成\n定义仍有未决项','8组附件清单 字段字典 来源快照 F I U证据清单'),
        ('3 质量与背景','已完成本轮','覆盖与缺测 QC交叉表 缺段及异常上下文 季节背景图'),
        ('4 第一问探索','当前位置\n持续推进','变量物理解释 非降水关系探索 配对样本量与年度差异'),
        ('5 目标与验证设计','待关键定义明确','灾害阶段 时区 数据可得时刻 研究范围 独立验证单位'),
        ('6 第二问评估','待开展','触发参数 临界条件 评估模型 时空可靠性'),
        ('7 第三问预报','待开展','触发时间区间 预警提前量 区间覆盖率 空间范围'),
        ('8 第四问仿真','待开展','地形与物源情景 参数表 过程可视化'),
        ('9 统一验证','待开展\n检查贯穿全程','基线对比 敏感性分析 事件或时空分组回测 复现实验'),
        ('10 论文与提交','素材持续积累\n定稿待开展','四问闭环 图表与引用 可复现附件 格式核验及提交')
    ],[1.32,1.18,4.36],highlight=3)
    paragraph(doc,'推进原则是先把数据和目标定义清楚，再建立稳定可解释的基线，最后只改进有明确问题依据的部分。模型复杂度、随机五折或深度学习框架均不作为预设答案。')

    # 第2页：已有事实和边界。
    next_page(doc,'已完成工作的证据')
    heading(doc,'工具与工程基础',2)
    paragraph(doc,'Miniforge的mc-base环境已用于实际审计和绘图，Git已记录本地版本。TeX Live和QGIS已准备；Origin及PyTorch按此前约定暂缓，当前探索无需为它们追加部署。[2]')
    heading(doc,'数据覆盖和质量发现',2)
    grid(doc,['对象','已核验结果','对后续工作的影响'],[
        ('小时观测','24份CSV 63个测量字段 37个QC字段','时间标签不等于完整有效观测，更不等于独立灾害样本'),
        ('2026年覆盖','声明窗内MET缺测低于0.01%\nRADM和SOIL为0% FLUX约3.99%','比例按测量单元格计算；数值齐全不代表质量合格'),
        ('连续缺段','FLUX有24字段共享4318个缺测小时标签','不能以简单插值伪造连续观测；缺段原因仍待核实'),
        ('范围告警','48条湿度告警对应QC为2\n含水率1.64对应QC为9','保存原值及邻近记录，不自动裁剪'),
        ('历史灾害','84个多边形 flag全部为1\n无明确逐事件发生时间','不能直接当作84个带时标正样本或构造可信负样本'),
        ('空间范围','DEM与23个历史面相交\n覆盖冰川掩膜7152个像元中心','面相交和完整覆盖分开统计；源区与县域样本范围不能混用'),
        ('目录差异','343个数值单元格位于官网起点之前','单列保存，不能自行判断是试运行或时间标注错误')
    ],[1.0,2.62,3.24])
    paragraph(doc,'以上来自项目审计与探索记录。[3][4][5] 官方小时观测时区为北京时间UTC+8，缺失码为9999.9；QC码0、1、2、4、8、9分别为正确、可疑、错误、提供方插补、缺失、未质控。[4]')
    heading(doc,'必须继续保留的解释边界',2)
    paragraph(doc,'站点距受灾口岸约65 km，官方限定其区域背景用途，不能直接代替灾区局地降水。Precip及FLUX共26个测量字段没有对应QC列，不能默认赋0；站点中英文纬度存在冲突。[4]')
    paragraph(doc,'原始附件未修改、没有新增填补。第一轮22项审计检查、元数据复核18项检查以及第二阶段6项边界测试和16项全量校验分别通过；这些验证的是代码、计数和证据一致性，不是模型预报能力。[3][4][5]')

    # 第3页：当前任务和真实进度。
    next_page(doc,'当前阶段与第一问推进')
    paragraph(doc,'赛题第一问要求分析边界层结构、天气过程、极端气象事件与泥石流的关联，并给出关联性最高的5个因素。[1] 当前开展的是这项任务的前置探索，变量之间的相关系数不能直接替代与灾害的关联排序。')
    heading(doc,'当前在做什么',2)
    paragraph(doc,'先依据表头和官方信息整理温度、湿度、风、气压、辐射与土壤变量。关系分析采用2022至2025完整历史年份，按北京时间标签日制作有覆盖门槛的描述性统计，保存配对样本量和跨年变化。QC为0与包含提供方插补记录的口径分别展示，作为敏感性观察，不预设最终QC准入规则。[6]')
    grid(doc,['变量组','可形成的探索内容','当前边界'],[
        ('温度与湿度','多高度记录 共同季节变化\n日内变化及变量配对关系','站点现象不直接代表灾区局地条件'),
        ('风与气压','高度间差异 天气背景\n与温湿变化的共同变化','梯度或稳定度指标须单独声明公式与假设'),
        ('辐射','短波和长波记录的背景变化\n与温度及土壤状态的关系','完整日统计仅为标签日描述，不是业务可得预报'),
        ('土壤','温度和含水率的共同变化\n不同深度记录的差异','不能直接代表冰川源区松散物或水分状态'),
        ('湍流与降水','保留FLUX原记录及缺段信息\n保留逐条降水曲线','FLUX无QC需另评估；累计降水方式尚未明确')
    ],[1.0,2.8,3.06])
    paragraph(doc,'本轮初步结果：QC为0的2米与20米气温日中位数高度同步，Spearman系数约0.999，配对1461天；2米气温与向下长波辐射的系数由0.913降至年内月份中心化后的0.390。关系对统计口径敏感，不能据此认定灾害驱动因素。[6]',bold_lead='本轮初步结果：')
    heading(doc,'第一问尚未完成的部分',2)
    paragraph(doc,'还要把多高度结构、天气过程与极端背景放到明确的比较框架中，区分季节共变、同次过程重复窗口和真正可检验的灾害关联。前5因素需要研究对象、对照依据和稳定性证据，不能按变量互相关的大小直接挑选。')
    paragraph(doc,'FLUX暂不进入有统一QC条件的主要相关分析，并不意味着从第一问永久删除。其补充分析须展示缺段与质量证据。降水累计、前期湿润指标和最终灾害时间标签继续暂缓。[4][6][7]')
    heading(doc,'本阶段完成标志',2)
    paragraph(doc,'形成变量解释表、可审查的关系图、每个系数的配对样本量及口径比较，并明确哪些结果只说明区域背景。这些产出用于选择下一轮假设，不能把第一问整体标记为已完成。')

    # 第4页：触发与预报路线。
    next_page(doc,'触发评估与预报的后续路线')
    heading(doc,'先明确目标和时间定义',2)
    paragraph(doc,'赛题案例写有2026年8月26日10时30分，但未明确它的时区、地点和灾害阶段。项目必须分别记录灾害链起始、泥石流触发、到达断面或口岸、首次报告等时间，保留各自来源、精度及转换证据，不能自动视作同一时刻。[1][4][7]')
    paragraph(doc,'同时核实小时标签代表的区间、发布延迟、降水增量或计数器及清零方式。研究源区、传播通道、历史清单和观测站的位置与覆盖须分别定义。关键定义不能确认时，保留条件性情景，不能写成已验证预测。')
    heading(doc,'第二问触发条件与评估模型',2)
    paragraph(doc,'目标是把地形、可动用物源、水源及触发条件组织成可解释的评估框架。先明确参数含义和单位，检查现有数据能支持哪些量，再按缺口补充有来源的卫星遥感、天气预报或过程记录。[1]')
    paragraph(doc,'采用能完整求解的成熟基线，先固定研究范围、输入和评价方式，再比较必要改进。临界条件应由物理约束、可核验事件或有说明的情景支持；在历史清单没有可靠发生时间、负样本缺少依据时，不伪造阈值的经验准确率。')
    paragraph(doc,'预期产出为参数字典、触发条件说明、基线结果、时间和空间尺度上的可靠性讨论。完成标志是输入与目标可追溯、计算可复现，且能够解释哪些条件下模型适用或不适用。')
    heading(doc,'第三问触发时间区间与预警',2)
    paragraph(doc,'目标是给出赛题要求的触发时间95%区间、预警时间、覆盖率与空间距离范围。[1] 先区分触发时刻不确定性、传播到达时刻和信息发布时间，确保每个预警截断点只使用当时可得数据。')
    paragraph(doc,'验证划分按独立事件及时间或空间单位设计，避免相邻小时、同次灾害重复窗口跨集合。区间方法必须与目标相符，记录名义95%与实测覆盖率的区别；区间宽度、覆盖率、预警提前量和空间范围分别报告。')
    paragraph(doc,'预期产出为时间基准表、回测设计、预报区间和误差记录。缺少独立带时标事件时，最多给出明确假设下的示例或情景，不宣称已经验证95%覆盖率。')

    # 第5页：仿真和统一验证。
    next_page(doc,'仿真与统一验证')
    heading(doc,'第四问仿真环境',2)
    paragraph(doc,'先在DEM确实覆盖的源区和通道内核验地形、投影、垂直基准和边界，再整理物源量、释放位置与时刻、水量、流变或阻力等参数。DEM文件名和裁剪日期不能证明其采集时间或灾前状态。[1][3][4]')
    paragraph(doc,'从参数少、含义清楚且能解释空间传播的方案起步，先测试守恒、稳定性与网格敏感性，再增加必要过程。具体求解器和物理模型留待地形、参数与校验资料明确后选择，不预设复杂度。')
    paragraph(doc,'输出包含参数表、初始与边界条件、过程快照及可视化、淹没或影响范围和数值诊断。情景仿真与经观测校准的灾害复现必须分开命名；没有过程观测时，不能把画面相似当作真实过程验证。')
    heading(doc,'贯穿四问的验证',2)
    grid(doc,['验证层面','需要检查的内容','保留的证据'],[
        ('数据与单位','原包校验 时间范围 投影与单位\n来源日期和可得时刻','数据清单 字段字典 原始哈希 元数据快照'),
        ('计算与实现','计数守恒 边界值 缺测与QC分离\n物理量转换和参数传递','边界测试 中间结果 运行日志'),
        ('统计与泛化','独立划分 基线对照\n跨年或跨区变化 样本不平衡','划分文件 配对样本量 逐次误差与区间结果'),
        ('物理与情景','量纲 守恒 稳定性\n网格与参数敏感性','参数范围 数值诊断 情景对照图'),
        ('可复现性','环境版本 随机种子 超参数\n数据版本和一键运行顺序','环境锁定 配置文件 脚本入口 结果清单')
    ],[1.05,2.85,2.96])
    heading(doc,'改进和停止条件',2)
    paragraph(doc,'每次改进说明要解决的具体问题，例如缺段导致样本覆盖不足、类别不平衡或空间外推失稳。只改变必要部分，与固定基线比较，并保留消融或敏感性结果；提升幅度不足以抵消不稳定性时，优先采用更稳定、可解释的结果。')
    paragraph(doc,'先冻结评价对象和验证集合，再做有限调参；新试验必须有可检验的问题。可重复获得的稳定结果优先，预留论文整理和最终提交核对时间。')

    # 第6页：论文、提交、近期行动及来源。
    next_page(doc,'论文提交与近期行动')
    heading(doc,'论文与附件形成过程',2)
    paragraph(doc,'从现在开始同步积累问题重述、来源与假设、变量解释、方法依据、验证结果和局限。四问各自形成输入、方法、结果和适用边界，最后通过同一研究范围、时间体系与参数表串联，避免各问口径不一致。')
    paragraph(doc,'图表标注单位、时段、样本量及不确定性含义，保存可编辑源图。论文定稿时核对摘要与正文结论一致、引用可定位、编号和交叉引用完整、模型与代码能对应；正式格式与截止时间以当届最新官方要求为准，目前不填写未经核验的日期或页数。')
    paragraph(doc,'提交前按官方要求核对命名、文件类型、页数、图表、引用和附件；用干净环境复跑关键结果，确认文件能打开，再完成提交及回执保存。建议HCX处理账号登录与最终提交确认，助手继续负责整理、计算、文档和核验；关键科学定义共同讨论。')
    heading(doc,'近期工作顺序',2)
    grid(doc,['顺序','近期动作','进入下一步的条件'],[
        ('1','完成第一问初步变量解释\n关系探索与口径比较','图表有样本量 有年度变化\n不误作灾害前5因素'),
        ('2','核实关键时间与降水定义\n确定研究对象和验证单位','证据能对应任务目标\n未决项仍明确保留'),
        ('3','建立可解释基线\n并准备地形与情景参数','输入和评价可追溯\n有校验方法再比较改进')
    ],[.6,3.0,3.26])
    paragraph(doc,'当前无需新增部署Origin或PyTorch，也无需为本轮分析额外登录网站。后续若官方资料需要账号访问，会明确提供入口和需要取得的内容。未确认信息继续按事实F、推断I、待核实U分开记录。')
    heading(doc,'依据与项目文件',2)
    references=[
        '[1] 2026年MathorCup大数据挑战赛赛道B初赛PDF，第1至2页。',
        '[2] docs/environment-setup.md；docs/git-usage.md；项目环境锁定文件。',
        '[3] docs/data-audit.md；docs/spatial-audit.md。',
        '[4] docs/metadata-review.md；config/analysis-constraints.yml。',
        '[5] docs/observation-exploration.md；output/observation_exploration/verification.json。',
        '[6] docs/variable-relations.md；output/variable_relations/，为本轮初步探索材料。',
        '[7] docs/pending-questions.md，保留降水、时刻、站点及派生数据处理等未决项。',
    ]
    for ref in references:
        p=paragraph(doc,ref)
        p.paragraph_format.space_after=Pt(3)
        for run in p.runs:
            font(run,9.5)
    OUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUT)
    print(json.dumps({'docx':str(OUT),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)},ensure_ascii=False))


if __name__=='__main__':
    main()
